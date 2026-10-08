"""
Export transcripts to DOCX and PDF (post-ASR only).
"""
from __future__ import annotations

import io
from typing import Any, Dict, List, Optional


def _safe_pdf_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bytearray):
        value = bytes(value)
    if isinstance(value, bytes):
        value = value.decode("utf-8", "replace")
    return str(value).encode("latin-1", "replace").decode("latin-1")


def build_docx(transcript: Dict[str, Any], extras: Optional[Dict[str, Any]] = None) -> bytes:
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    title = transcript.get("title") or "Lecture Transcript"
    doc.add_heading(title, 0)

    meta = []
    if transcript.get("course_code"):
        meta.append(f"Course: {transcript['course_code']}")
    if transcript.get("speaker"):
        meta.append(f"Speaker: {transcript['speaker']}")
    if transcript.get("duration"):
        meta.append(f"Duration: {transcript['duration']}")
    if transcript.get("decode_mode"):
        meta.append(f"Decoding: {transcript['decode_mode']}")
    if transcript.get("processed_at"):
        meta.append(f"Processed: {transcript['processed_at']}")
    if meta:
        doc.add_paragraph("\n".join(meta))

    extras = extras or {}
    if extras.get("ai_summary"):
        doc.add_heading("AI Summary", level=1)
        doc.add_paragraph(extras["ai_summary"])

    doc.add_heading("Transcript", level=1)
    segments: List[Dict[str, Any]] = transcript.get("segments") or []
    if segments:
        for seg in segments:
            ts = seg.get("timestamp", "")
            text = seg.get("text", "")
            para = doc.add_paragraph()
            run = para.add_run(f"[{ts}] ")
            run.bold = True
            para.add_run(text)
    else:
        doc.add_paragraph(transcript.get("full_text") or "")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def build_pdf(transcript: Dict[str, Any], extras: Optional[Dict[str, Any]] = None) -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _safe_pdf_text(transcript.get("title") or "Lecture Transcript"))
    pdf.ln(4)

    pdf.set_font("Helvetica", "", 10)
    meta_lines = []
    for key, label in (
        ("course_code", "Course"),
        ("speaker", "Speaker"),
        ("duration", "Duration"),
        ("decode_mode", "Decoding"),
        ("processed_at", "Processed"),
    ):
        if transcript.get(key):
            meta_lines.append(_safe_pdf_text(f"{label}: {transcript[key]}"))
    if meta_lines:
        pdf.multi_cell(0, 6, "\n".join(meta_lines))
        pdf.ln(4)

    extras = extras or {}
    if extras.get("ai_summary"):
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, "AI Summary", ln=True)
        pdf.set_font("Helvetica", "", 10)
        pdf.multi_cell(0, 6, _safe_pdf_text(extras["ai_summary"]))
        pdf.ln(4)

    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "Transcript", ln=True)
    pdf.set_font("Helvetica", "", 10)

    segments: List[Dict[str, Any]] = transcript.get("segments") or []
    if segments:
        for seg in segments:
            ts = _safe_pdf_text(seg.get("timestamp", ""))
            text = _safe_pdf_text(seg.get("text", ""))
            line = f"[{ts}] {text}"
            pdf.multi_cell(0, 6, line)
            pdf.ln(1)
    else:
        text = _safe_pdf_text(transcript.get("full_text") or "")
        pdf.multi_cell(0, 6, text)

    raw_pdf = pdf.output(dest="S")
    if isinstance(raw_pdf, bytearray):
        pdf_bytes = bytes(raw_pdf)
    elif isinstance(raw_pdf, bytes):
        pdf_bytes = raw_pdf
    else:
        pdf_bytes = str(raw_pdf).encode("latin-1", "replace")
    return pdf_bytes
